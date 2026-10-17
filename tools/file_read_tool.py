from pathlib import Path
from typing import Annotated

from langchain_core.tools import tool

from api.context import get_session_context
from api.monitor import monitor
from utils.path_utils import resolve_in_session

# 上下文保护：一个几十页的 PDF 全塞给模型，会直接撑爆上下文窗口。
MAX_CHARS = 20000

_TEXT_SUFFIXES = {".md", ".markdown", ".txt", ".csv", ".json", ".log", ".yaml", ".yml"}


def _read_text(path: Path) -> str:
    try:
        return path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        # 非 UTF-8 文本（例如 GBK 导出的 CSV）尽量读，不要直接失败
        return path.read_text(encoding="utf-8", errors="replace")


def _read_docx(path: Path) -> str:
    try:
        import docx
    except ImportError:
        return "读取失败：未安装 python-docx，无法读取 .docx（pip install python-docx）"

    document = docx.Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs]
    # 表格内容不在 paragraphs 里，需要单独提取
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _read_pdf(path: Path) -> str:
    try:
        import pypdf
    except ImportError:
        return "读取失败：未安装 pypdf，无法读取 .pdf（pip install pypdf）"

    reader = pypdf.PdfReader(str(path))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _read_excel(path: Path) -> str:
    try:
        import pandas as pd
    except ImportError:
        return "读取失败：未安装 pandas，无法读取 Excel（pip install pandas openpyxl）"

    sheets = pd.read_excel(str(path), sheet_name=None)
    parts = []
    for name, frame in sheets.items():
        parts.append(f"【工作表：{name}】")
        parts.append(f"行数：{len(frame)}，列数：{len(frame.columns)}")
        parts.append(f"列名：{', '.join(str(column) for column in frame.columns)}")
        parts.append("[前 10 行]")
        parts.append(frame.head(10).to_string(index=False))
        parts.append("")
    return "\n".join(parts)


def _truncate(text: str, filename: str) -> str:
    if len(text) <= MAX_CHARS:
        return text
    return (
        text[:MAX_CHARS]
        + f"\n\n[已截断：{filename} 共 {len(text)} 字符，"
        f"仅返回前 {MAX_CHARS} 字符]"
    )


@monitor.traced_tool("文件内容读取工具")
@tool
def read_file_content(
    filename: Annotated[
        str,
        "要读取的文件名或相对路径，例如 report.md 或 资料/数据.xlsx。"
        "支持 .md/.txt/.csv/.json/.docx/.pdf/.xlsx/.xls",
    ],
) -> str:
    """读取当前会话目录里的文件内容（用户上传的附件、之前生成的报告等）。

    读取结果超过 20000 字符时会截断，并明确标注。
    """
    try:
        session_dir = get_session_context()
        if not session_dir:
            return "读取失败：当前没有设置会话目录（session_dir 未初始化）。"

        path = resolve_in_session(session_dir, filename)
        if path is None:
            return "读取失败：filename 不能为空。"
        if not path.is_file():
            return f"读取失败：文件不存在 {path.name}"

        suffix = path.suffix.lower()
        if suffix in _TEXT_SUFFIXES:
            text = _read_text(path)
        elif suffix == ".docx":
            text = _read_docx(path)
        elif suffix == ".pdf":
            text = _read_pdf(path)
        elif suffix in (".xlsx", ".xls"):
            text = _read_excel(path)
        else:
            supported = "、".join(
                sorted(_TEXT_SUFFIXES | {".docx", ".pdf", ".xlsx", ".xls"})
            )
            return f"读取失败：不支持的文件格式 {suffix or '（无扩展名）'}。支持：{supported}"

        if not text.strip():
            return (
                f"文件 {path.name} 没有可提取的文本内容"
                "（可能是扫描版 PDF 或空文件）。"
            )
        return _truncate(text, path.name)
    except Exception as error:
        return f"读取失败：{type(error).__name__}: {error}"
